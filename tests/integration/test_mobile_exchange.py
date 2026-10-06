"""Mobile import review survives restarts and keeps exact, reversible cash flow."""

import json
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from openledger.mobile.bridge import MobileLedger
from tests.integration.test_mobile_bridge import NOW, account, digest, invoke, record_body
from tests.integration.test_mobile_mutations import mutate


def source(mobile: MobileLedger, format: str = "csv") -> str:
    """Write synthetic fixtures only inside the application-owned staging directory."""
    rows = [
        ("日期", "金额", "类型", "备注", "external_source", "external_transaction_id"),
        ("2026-10-03", "12.34", "支出", "合成午餐", "synthetic", "one"),
        ("2026-10-04", "100.01", "收入", "合成工资", "synthetic", "two"),
    ]
    path = mobile.files.staging / f"sample.{format}"
    if format == "csv":
        import csv

        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(rows)
    else:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
    return path.name


def preview(mobile: MobileLedger, filename: str) -> dict[str, Any]:
    snapshot = invoke(mobile, "snapshot")["data"]
    mapping = {
        "book_id": snapshot["books"][0]["id"],
        "account_id": snapshot["accounts"][0]["id"],
        "income_category_id": next(
            row["id"] for row in snapshot["categories"] if row["transaction_kind"] == "income"
        ),
        "expense_category_id": next(
            row["id"] for row in snapshot["categories"] if row["transaction_kind"] == "expense"
        ),
    }
    result = invoke(mobile, "import_preview", {"filename": filename, "mapping": mapping})
    assert result["ok"], result
    return dict(result["data"])


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_review_commit_reopen_duplicate_and_revert(tmp_path: Path, format: str) -> None:
    directory = tmp_path / "ledger"
    mobile = MobileLedger(str(directory), clock=lambda: NOW)
    account(mobile)
    filename = source(mobile, format)
    before = digest(mobile)
    plan = preview(mobile, filename)
    assert digest(mobile) == before
    assert all(not row["issues"] for row in plan["rows"])
    mobile = MobileLedger(str(directory), clock=lambda: NOW)
    decision = invoke(
        mobile,
        "import_prepare",
        {"preview_id": plan["preview_id"], "selected_rows": [2, 3], "request_id": str(uuid4())},
    )
    assert decision["ok"], decision
    mobile = MobileLedger(str(directory), clock=lambda: NOW)
    committed = invoke(mobile, "import_commit", decision["data"])
    assert committed["ok"], committed
    assert invoke(mobile, "import_commit", decision["data"])["data"]["replayed"]
    assert invoke(mobile, "snapshot")["data"]["overview"]["total_assets_minor"] == 108767
    duplicates = preview(mobile, filename)
    assert all(row["issues"] for row in duplicates["rows"])
    batch = invoke(mobile, "snapshot")["data"]["import_batches"][0]
    assert mutate(
        mobile, "import.revert.v1", {"id": batch["id"], "expected_version": batch["version"]}
    )["ok"]
    assert invoke(mobile, "snapshot")["data"]["overview"]["total_assets_minor"] == 100000


def test_changed_source_and_decision_rejected_without_writes(tmp_path: Path) -> None:
    mobile = MobileLedger(str(tmp_path / "ledger"), clock=lambda: NOW)
    account(mobile)
    plan = preview(mobile, source(mobile))
    saved_source = mobile.files.directory / "imports" / f"source-{plan['preview_id']}.csv"
    saved_source.write_bytes(b"changed")
    before = digest(mobile)
    reply = invoke(
        mobile,
        "import_prepare",
        {"preview_id": plan["preview_id"], "selected_rows": [2], "request_id": str(uuid4())},
    )
    assert reply["error"]["code"] == "IMPORT_FILE_CHANGED"
    assert digest(mobile) == before


@pytest.mark.parametrize("corrupt", ["not JSON", "{}"])
def test_damaged_saved_plan_returns_safe_error(tmp_path: Path, corrupt: str) -> None:
    mobile = MobileLedger(str(tmp_path / "ledger"), clock=lambda: NOW)
    account(mobile)
    plan = preview(mobile, source(mobile))
    path = mobile.files.directory / "imports" / f"preview-{plan['preview_id']}.json"
    path.write_text(corrupt, encoding="utf-8")
    before = digest(mobile)
    reply = invoke(
        mobile,
        "import_prepare",
        {
            "preview_id": plan["preview_id"],
            "request_id": str(uuid4()),
            "selected_rows": [2],
        },
    )
    assert reply["error"]["code"] == "IMPORT_PREVIEW_CHANGED"
    assert digest(mobile) == before


@pytest.mark.parametrize("selected", [[True], [2, 2], [999], [], [2.0]])
def test_invalid_selection_cannot_write(tmp_path: Path, selected: list[object]) -> None:
    mobile = MobileLedger(str(tmp_path / "ledger"), clock=lambda: NOW)
    account(mobile)
    plan = preview(mobile, source(mobile))
    before = digest(mobile)
    reply = invoke(
        mobile,
        "import_prepare",
        {"preview_id": plan["preview_id"], "selected_rows": selected, "request_id": str(uuid4())},
    )
    assert not reply["ok"], reply
    assert digest(mobile) == before


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_and_report_exact_precision(tmp_path: Path, format: str) -> None:
    mobile = MobileLedger(str(tmp_path / "ledger"), clock=lambda: NOW)
    account(mobile)
    body = record_body(mobile)
    body["fields"]["amount"] = "0.01"
    body["fields"]["note"] = "=SUM(1,2)"
    assert invoke(mobile, "record", body)["ok"]
    result = invoke(
        mobile, "export_transactions", {"filename": "export." + format, "format": format}
    )
    assert result["ok"], result
    assert result["data"]["row_count"] == 2
    if format == "xlsx":
        workbook = load_workbook(mobile.files.staging / "export.xlsx")
        assert all(cell.data_type != "f" for row in workbook.active or [] for cell in row)
        workbook.close()
    report = invoke(mobile, "analytics", {"start_on": "2026-10-01", "end_on": "2026-10-04"})
    assert report["ok"], report
    assert report["data"]["totals"]["net_expense_minor"] == "1"
    assert report["data"]["categories"][0]["share"] == "1"
    assert report["data"]["months"][0]["totals"]["gross_expense_minor"] == "1"


def test_corrupt_confirmed_decision_cannot_be_committed(tmp_path: Path) -> None:
    mobile = MobileLedger(str(tmp_path / "ledger"), clock=lambda: NOW)
    account(mobile)
    plan = preview(mobile, source(mobile))
    decision = invoke(
        mobile,
        "import_prepare",
        {"preview_id": plan["preview_id"], "selected_rows": [2], "request_id": str(uuid4())},
    )["data"]
    path = mobile.files.directory / "imports" / f"decision-{decision['request_id']}.json"
    stored = json.loads(path.read_text(encoding="utf-8"))
    stored["rows"][0]["fields"]["amount_minor"] = 999
    path.write_text(json.dumps(stored), encoding="utf-8")
    before = digest(mobile)
    result = invoke(mobile, "import_commit", decision)
    assert result["error"]["code"] == "IMPORT_PREVIEW_CHANGED"
    assert digest(mobile) == before
