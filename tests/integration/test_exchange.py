"""Real bounded CSV/XLSX exchange, confirmed imports and reversible file publication."""

import csv
import zipfile
from dataclasses import replace
from datetime import UTC, date, datetime
from pathlib import Path
from typing import cast
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from openledger.application.dto.exchange import ExchangeMapping, ImportPreview
from openledger.application.dto.ledger import RefundFields, TransactionFields, TransferFields
from openledger.application.dto.queries import TransactionFilter
from openledger.domain.errors import LedgerError
from openledger.domain.money import MAX_EVENT_MINOR
from openledger.infrastructure import exchange
from openledger.infrastructure.database.database import Database
from openledger.infrastructure.exchange import ExchangeService, automatic_columns, read_table
from openledger.infrastructure.ledger import LedgerService

pytestmark = pytest.mark.integration
TODAY = date(2026, 10, 2)


def uid() -> str:
    return str(uuid4())


def refs(ledger: LedgerService) -> dict[str, str]:
    return {
        **{
            kind: str(
                next(row["id"] for row in ledger.entities("account") if row["account_type"] == kind)
            )
            for kind in ("cash", "bank")
        },
        "book": str(ledger.entities("book")[0]["id"]),
        **{
            kind: str(
                next(
                    row["id"]
                    for row in ledger.entities("category")
                    if row["transaction_kind"] == kind
                )
            )
            for kind in ("income", "expense")
        },
    }


def mapping(
    ledger: LedgerService,
    *,
    columns: tuple[tuple[str, str], ...] = (),
    sheet: str | None = None,
) -> ExchangeMapping:
    r = refs(ledger)
    return ExchangeMapping(
        columns=columns,
        sheet=sheet,
        book_id=r["book"],
        account_id=r["cash"],
        income_category_id=r["income"],
        expense_category_id=r["expense"],
        from_account_id=r["cash"],
        to_account_id=r["bank"],
    )


def csv_file(
    path: Path,
    headers: tuple[str, ...],
    rows: tuple[tuple[str, ...], ...],
    *,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
) -> Path:
    with path.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def xlsx_file(
    path: Path,
    headers: tuple[str, ...],
    rows: tuple[tuple[object, ...], ...],
    *,
    sheet_name: str = "账单",
) -> Path:
    workbook = Workbook()
    try:
        sheet = cast(Worksheet, workbook.active)
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        workbook.save(path)
    finally:
        workbook.close()
    return path


def preview_csv(
    ledger: LedgerService,
    path: Path,
    rows: tuple[tuple[str, ...], ...] = (("支出", "25.80", "2026-10-02", "咖啡"),),
    *,
    headers: tuple[str, ...] = ("类型", "金额", "日期", "备注"),
    plan: ExchangeMapping | None = None,
) -> ImportPreview:
    csv_file(path, headers, rows)
    chosen = plan or mapping(ledger)
    return ExchangeService(ledger).preview(read_table(path, chosen), chosen)


def counts(ledger: LedgerService) -> tuple[int, ...]:
    with ledger.database.read() as connection:
        return tuple(
            int(connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])
            for table in (
                "transactions",
                "account_entries",
                "import_batches",
                "audit_events",
                "change_log",
                "command_receipts",
            )
        )


def record(
    ledger: LedgerService,
    *,
    amount: int = 2_580,
    kind: str = "expense",
    occurred_on: date = TODAY,
    time_zone: str = "UTC",
    occurrence_precision: str = "date",
    occurred_at_utc: datetime | None = None,
    note: str | None = None,
    source: str = "manual",
    source_text: str | None = None,
    tag_ids: tuple[str, ...] = (),
    counterparty: str | None = None,
    merchant: str | None = None,
    location: str | None = None,
) -> str:
    r = refs(ledger)
    fields = TransactionFields(
        kind,
        amount,
        r["cash"],
        r["book"],
        r[kind],
        occurred_on,
        time_zone=time_zone,
        occurrence_precision=occurrence_precision,
        occurred_at_utc=occurred_at_utc,
        note=note,
        source=source,
        source_text=source_text,
        tag_ids=tag_ids,
        counterparty=counterparty,
        merchant=merchant,
        location=location,
    )
    identifier = uid()
    ledger.record(fields, request_id=uid(), transaction_id=identifier)
    return identifier


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-8", "gb18030"])
@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_csv_encodings_and_delimiters_keep_chinese_cells(
    tmp_path: Path, encoding: str, delimiter: str
) -> None:
    path = csv_file(
        tmp_path / "中文账单.csv",
        ("日期", "金额", "备注"),
        (("2026-10-02", "25.80", "咖啡，朋友\n聚会"),),
        encoding=encoding,
        delimiter=delimiter,
    )
    table = read_table(path, ExchangeMapping(encoding=encoding))
    assert table.headers == ("日期", "金额", "备注")
    assert table.rows == (("2026-10-02", "25.80", "咖啡，朋友\n聚会"),)
    assert table.source_format == "csv" and len(table.file_digest) == 64


def test_chinese_automatic_and_explicit_mapping_do_not_write(
    ledger: LedgerService, tmp_path: Path
) -> None:
    before_counts, before_balances = counts(ledger), ledger.balances()
    preview = preview_csv(ledger, tmp_path / "账单.csv")
    assert not preview.rows[0].issues
    fields = cast(dict[str, object], preview.rows[0].fields)
    assert fields["kind"] == "expense" and fields["amount_minor"] == 2_580
    assert fields["note"] == "咖啡" and fields["time_zone"] == "Asia/Shanghai"
    assert dict(automatic_columns(("收支类型", "金额(元)", "交易日期"))) == {
        "kind": "收支类型",
        "amount": "金额(元)",
        "occurred_on": "交易日期",
    }
    explicit = preview_csv(
        ledger,
        tmp_path / "自定义.csv",
        (("2026-10-02", "100.01"),),
        headers=("记账日", "花销"),
        plan=mapping(ledger, columns=(("occurred_on", "记账日"), ("amount", "花销"))),
    )
    assert cast(dict[str, object], explicit.rows[0].fields)["amount_minor"] == 10_001
    assert counts(ledger) == before_counts and ledger.balances() == before_balances


def test_preview_commit_and_revert_share_real_financial_boundary(
    ledger: LedgerService, tmp_path: Path
) -> None:
    service = ExchangeService(ledger)
    preview = preview_csv(ledger, tmp_path / "导入.csv")
    payload = service.commit_payload(preview, [2])
    result = ledger.execute(uid(), "import.commit.v1", payload)
    assert result.data["accepted_row_count"] == 1
    assert ledger.total_assets() == 97_420
    detail = ledger.transaction(preview.rows[0].transaction_id)
    assert detail["note"] == "咖啡" and detail["import_source_row"] == 2
    ledger.execute(uid(), "import.revert.v1", {"id": preview.batch_id, "expected_version": 1})
    assert ledger.total_assets() == 100_000
    assert ledger.import_batches()[0]["status"] == "reverted"


@pytest.mark.parametrize(
    "amount,issue",
    [
        ("1.001", "AMOUNT_PRECISION"),
        ("1e3", "INVALID_AMOUNT"),
        ("-1", "INVALID_AMOUNT"),
        ("0", "INVALID_AMOUNT"),
        ("NaN", "INVALID_AMOUNT"),
        ("1,000", "INVALID_AMOUNT"),
        ("1000000000000", "AMOUNT_OUT_OF_RANGE"),
    ],
)
def test_invalid_amount_is_a_blocked_preview_row(
    ledger: LedgerService, tmp_path: Path, amount: str, issue: str
) -> None:
    preview = preview_csv(ledger, tmp_path / "bad.csv", (("支出", amount, "2026-10-02", ""),))
    assert issue in preview.rows[0].issues and preview.rows[0].fields is None
    with pytest.raises(LedgerError, match="IMPORT_SELECTION_INVALID"):
        ExchangeService(ledger).commit_payload(preview, [2])
    assert not ledger.import_batches()


@pytest.mark.parametrize(
    "day,issue",
    [
        ("2026-10-03", "FUTURE_DATE"),
        ("2026-02-30", "INVALID_DATE"),
        ("2025-12-31", "BEFORE_ACCOUNT_START"),
    ],
)
def test_bad_dates_are_blocked_before_commit(
    ledger: LedgerService, tmp_path: Path, day: str, issue: str
) -> None:
    preview = preview_csv(ledger, tmp_path / "date.csv", (("支出", "25", day, ""),))
    assert issue in preview.rows[0].issues
    assert ledger.total_assets() == 100_000


def test_archived_account_is_visible_as_blocked_reference(
    ledger: LedgerService, tmp_path: Path
) -> None:
    r = refs(ledger)
    ledger.execute(
        uid(),
        "account.archive.v1",
        {
            "id": r["cash"],
            "expected_version": 1,
            "archived": True,
            "replacement_default_id": r["bank"],
        },
    )
    chosen = ExchangeMapping(
        book_id=r["book"], account_id=r["cash"], expense_category_id=r["expense"]
    )
    preview = preview_csv(ledger, tmp_path / "archived.csv", plan=chosen)
    assert "ENTITY_ARCHIVED" in preview.rows[0].issues


def test_unknown_explicit_account_does_not_silently_use_default(
    ledger: LedgerService, tmp_path: Path
) -> None:
    preview = preview_csv(
        ledger,
        tmp_path / "unknown-account.csv",
        (("expense", "25", "2026-10-02", "不存在的账户"),),
        headers=("kind", "amount", "occurred_on", "account_name"),
    )
    assert preview.rows[0].fields is None
    assert "IMPORT_REFERENCE_REQUIRED" in preview.rows[0].issues


def test_conflicting_explicit_id_and_name_block_account_routing(
    ledger: LedgerService, tmp_path: Path
) -> None:
    r = refs(ledger)
    bank = next(account for account in ledger.entities("account") if account["id"] == r["bank"])
    preview = preview_csv(
        ledger,
        tmp_path / "conflicting-account.csv",
        (("expense", "25", "2026-10-02", r["cash"], str(bank["name"])),),
        headers=("kind", "amount", "occurred_on", "account_id", "account_name"),
    )
    assert preview.rows[0].fields is None
    assert "IMPORT_REFERENCE_CONFLICT" in preview.rows[0].issues


@pytest.mark.parametrize(
    "columns",
    [
        (("amount", "金额"), ("amount", "日期")),
        (("amount", "不存在的列"),),
    ],
)
def test_bad_header_mapping_fails_before_preview(
    ledger: LedgerService, tmp_path: Path, columns: tuple[tuple[str, str], ...]
) -> None:
    path = csv_file(tmp_path / "mapping.csv", ("金额", "日期"), (("25", "2026-10-02"),))
    with pytest.raises(LedgerError, match="IMPORT_MAPPING_INVALID"):
        ExchangeService(ledger).preview(read_table(path), mapping(ledger, columns=columns))


def test_xlsx_sheet_dates_and_formula_rows_are_explicit(
    ledger: LedgerService, tmp_path: Path
) -> None:
    path = xlsx_file(
        tmp_path / "中文.xlsx",
        ("类型", "金额", "日期", "备注"),
        (
            ("支出", "25.80", datetime(2026, 10, 2), "正常"),
            ("支出", "=10+15", datetime(2026, 10, 2), "公式"),
        ),
    )
    chosen = mapping(ledger, sheet="账单")
    table = read_table(path, chosen)
    assert table.sheets == ("账单",) and table.formula_rows == frozenset({3})
    assert table.rows[0][2] == "2026-10-02"
    preview = ExchangeService(ledger).preview(table, chosen)
    assert not preview.rows[0].issues
    assert "IMPORT_FORMULA_FORBIDDEN" in preview.rows[1].issues
    with pytest.raises(LedgerError, match="IMPORT_INVALID_FILE"):
        read_table(path, replace(chosen, sheet="不存在"))


@pytest.mark.parametrize(
    "headers,rows,issue",
    [
        (("", "金额"), (("日期", "25"),), "IMPORT_HEADER_INVALID"),
        (("金额", "金额"), (("25", "25"),), "IMPORT_HEADER_INVALID"),
        (("金额", "日期"), (), "IMPORT_EMPTY_FILE"),
        (("金额", "日期"), (("25",),), "IMPORT_COLUMN_COUNT"),
    ],
)
def test_invalid_headers_and_width_fail_the_file(
    tmp_path: Path, headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...], issue: str
) -> None:
    path = csv_file(tmp_path / "header.csv", headers, rows)
    with pytest.raises(LedgerError, match=issue):
        read_table(path)


def test_empty_file_and_formula_header_are_rejected(tmp_path: Path) -> None:
    empty = tmp_path / "empty.csv"
    empty.write_bytes(b"")
    with pytest.raises(LedgerError, match="IMPORT_HEADER_INVALID"):
        read_table(empty)
    formula = xlsx_file(tmp_path / "formula.xlsx", ("=1+1", "金额"), (("日期", "25"),))
    with pytest.raises(LedgerError, match="IMPORT_HEADER_INVALID"):
        read_table(formula)


@pytest.mark.parametrize(
    "name", ["../evil.xml", "/absolute.xml", "C:/drive.xml", "folder\\evil.xml", "a/../b.xml"]
)
def test_xlsx_unsafe_zip_names_are_rejected_before_loading(tmp_path: Path, name: str) -> None:
    path = tmp_path / "unsafe.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(name, "malformed worksheet")
    if "\\" in name:
        # Windows ZipInfo normalizes separators while writing; reproduce stored raw bytes.
        path.write_bytes(path.read_bytes().replace(name.replace("\\", "/").encode(), name.encode()))
    with pytest.raises(LedgerError, match="IMPORT_ARCHIVE_UNSAFE"):
        read_table(path)


def test_xlsx_duplicate_member_and_compression_bomb_are_rejected(tmp_path: Path) -> None:
    duplicate = tmp_path / "duplicate.xlsx"
    with zipfile.ZipFile(duplicate, "w") as archive:
        archive.writestr("sheet.xml", "one")
        with pytest.warns(UserWarning, match="Duplicate name"):
            archive.writestr("sheet.xml", "two")
    with pytest.raises(LedgerError, match="IMPORT_ARCHIVE_UNSAFE"):
        read_table(duplicate)
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("sheet.xml", "0" * 200_000)
    with pytest.raises(LedgerError, match="IMPORT_ARCHIVE_UNSAFE"):
        read_table(bomb)


@pytest.mark.parametrize("limit", ["rows", "columns", "bytes", "cell"])
@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_read_limits_are_enforced(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, limit: str, format: str
) -> None:
    headers, rows = ("a", "b", "c"), (("1", "2", "3"),) * 3
    path = tmp_path / ("large." + format)
    if format == "csv":
        csv_file(path, headers, rows)
    else:
        xlsx_file(path, headers, rows)
    setting, value = {
        "rows": ("MAX_ROWS", 2),
        "columns": ("MAX_COLUMNS", 2),
        "bytes": ("MAX_FILE_BYTES", 8),
        "cell": ("MAX_CELL", 0),
    }[limit]
    monkeypatch.setattr(exchange, setting, value)
    with pytest.raises(LedgerError):
        read_table(path)


def test_read_and_preview_cancellation_preserve_the_database(
    ledger: LedgerService, tmp_path: Path
) -> None:
    path = csv_file(tmp_path / "cancel.csv", ("金额", "日期"), (("25", "2026-10-02"),))
    before = counts(ledger)
    with pytest.raises(LedgerError, match="EXCHANGE_CANCELLED"):
        read_table(path, cancel=lambda: True)
    with pytest.raises(LedgerError, match="EXCHANGE_CANCELLED"):
        ExchangeService(ledger).preview(read_table(path), mapping(ledger), cancel=lambda: True)
    assert counts(ledger) == before


def test_exact_duplicates_remain_blocked_after_revert(
    ledger: LedgerService, tmp_path: Path
) -> None:
    path = tmp_path / "duplicates.csv"
    preview = preview_csv(ledger, path)
    service = ExchangeService(ledger)
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(preview, [2]))
    for _attempt in range(2):
        again = service.preview(read_table(path), mapping(ledger))
        assert again.rows[0].exact_duplicate
        assert "DUPLICATE_IMPORT" in again.rows[0].issues
        with pytest.raises(LedgerError, match="IMPORT_SELECTION_INVALID"):
            service.commit_payload(again, [2])
        if _attempt == 0:
            ledger.execute(
                uid(), "import.revert.v1", {"id": preview.batch_id, "expected_version": 1}
            )


def test_possible_duplicates_require_visible_selection_but_are_not_dropped(
    ledger: LedgerService, tmp_path: Path
) -> None:
    existing = record(ledger)
    preview = preview_csv(
        ledger, tmp_path / "fuzzy.csv", (("支出", "25.80", "2026-10-02", ""),) * 2
    )
    assert all(existing in row.possible_duplicates for row in preview.rows)
    assert preview.rows[0].transaction_id in preview.rows[1].possible_duplicates
    assert all(not row.exact_duplicate and not row.issues for row in preview.rows)
    service = ExchangeService(ledger)
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(preview, [3]))
    assert ledger.total_assets() == 94_840
    assert ledger.import_batches()[0]["accepted_row_count"] == 1


def test_external_duplicates_in_file_are_blocked(ledger: LedgerService, tmp_path: Path) -> None:
    chosen = mapping(ledger)
    path = csv_file(
        tmp_path / "external.csv",
        ("kind", "amount", "occurred_on", "external_source", "external_transaction_id"),
        (("expense", "25", "2026-10-02", "bank:test", "order-one"),) * 2,
    )
    preview = ExchangeService(ledger).preview(read_table(path), chosen)
    assert not preview.rows[0].issues
    assert preview.rows[1].exact_duplicate and "DUPLICATE_IMPORT" in preview.rows[1].issues


def test_file_changed_and_invalid_selection_are_rejected(
    ledger: LedgerService, tmp_path: Path
) -> None:
    path = tmp_path / "changed.csv"
    preview = preview_csv(ledger, path)
    service = ExchangeService(ledger)
    for selection in ([], [999]):
        with pytest.raises(LedgerError, match="IMPORT_SELECTION_INVALID"):
            service.commit_payload(preview, selection)
    path.write_text("金额,日期\n999,2026-10-02\n", encoding="utf-8")
    with pytest.raises(LedgerError, match="IMPORT_FILE_CHANGED"):
        service.commit_payload(preview, [2])
    assert not ledger.import_batches()


def test_mutated_preview_fields_cannot_publish_a_different_displayed_intent(
    ledger: LedgerService, tmp_path: Path
) -> None:
    preview = preview_csv(ledger, tmp_path / "preview-mutated.csv")
    assert preview.plan_digest
    cast(dict[str, object], preview.rows[0].fields)["amount_minor"] = 1
    with pytest.raises(LedgerError, match="IMPORT_PREVIEW_CHANGED"):
        ExchangeService(ledger).commit_payload(preview, [2])
    assert not ledger.import_batches() and ledger.total_assets() == 100_000


def test_refund_selection_requires_its_new_expense(ledger: LedgerService, tmp_path: Path) -> None:
    original, refund = uid(), uid()
    path = csv_file(
        tmp_path / "dependency.csv",
        ("transaction_id", "kind", "amount", "occurred_on", "original_transaction_id"),
        (
            (refund, "expense_refund", "2", "2026-10-02", original),
            (original, "expense", "25", "2026-10-02", ""),
        ),
    )
    service = ExchangeService(ledger)
    preview = service.preview(read_table(path), mapping(ledger))
    assert all(not row.issues for row in preview.rows)
    with pytest.raises(LedgerError, match="IMPORT_DEPENDENCY_REQUIRED"):
        service.commit_payload(preview, [2])
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(preview, [2, 3]))
    assert ledger.total_assets() == 97_700


def test_refund_dependency_uses_first_eligible_duplicate_source_expense(
    ledger: LedgerService, tmp_path: Path
) -> None:
    original = uid()
    path = csv_file(
        tmp_path / "duplicate-source-refund.csv",
        ("transaction_id", "kind", "amount_minor", "occurred_on", "original_transaction_id"),
        (
            (original, "expense", "10000", "2026-10-02", ""),
            (original, "expense", "10000", "2026-10-02", ""),
            (uid(), "expense_refund", "1000", "2026-10-02", original),
        ),
    )
    service = ExchangeService(ledger)
    preview = service.preview(read_table(path), mapping(ledger))
    first, duplicate, refund = preview.rows
    assert first.source_row_number == 2 and not first.issues
    assert duplicate.source_row_number == 3 and duplicate.exact_duplicate
    assert "DUPLICATE_IMPORT" in duplicate.issues
    assert refund.source_row_number == 4 and not refund.issues
    assert cast(dict[str, object], refund.fields)["original_transaction_id"] == first.transaction_id
    assert (
        cast(dict[str, object], refund.fields)["original_transaction_id"]
        != duplicate.transaction_id
    )
    payload = service.commit_payload(preview, [2, 4])
    result = ledger.execute(uid(), "import.commit.v1", payload)
    assert result.data["accepted_row_count"] == 2
    assert ledger.total_assets() == 91_000
    imported_refund = ledger.transaction(refund.transaction_id)
    assert imported_refund["original_transaction_id"] == first.transaction_id
    assert ledger.import_batches()[0]["accepted_row_count"] == 2


def test_invalid_first_source_row_rebinds_refund_to_eligible_expense(
    ledger: LedgerService, tmp_path: Path
) -> None:
    original = uid()
    path = csv_file(
        tmp_path / "invalid-first-original.csv",
        ("transaction_id", "kind", "amount_minor", "occurred_on", "original_transaction_id"),
        (
            (original, "expense", "bad", "2026-10-02", ""),
            (original, "expense", "10000", "2026-10-02", ""),
            (uid(), "expense_refund", "1000", "2026-10-02", original),
        ),
    )
    service = ExchangeService(ledger)
    preview = service.preview(read_table(path), mapping(ledger))
    invalid, expense, refund = preview.rows
    assert "INVALID_AMOUNT" in invalid.issues and invalid.fields is None
    assert not expense.issues and not refund.issues
    assert (
        cast(dict[str, object], refund.fields)["original_transaction_id"] == expense.transaction_id
    )
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(preview, [3, 4]))
    assert ledger.total_assets() == 91_000
    assert (
        ledger.transaction(refund.transaction_id)["original_transaction_id"]
        == expense.transaction_id
    )


@pytest.mark.parametrize("first_kind,second_amount", [("income", "10000"), ("expense", "20000")])
def test_conflicting_source_identity_blocks_every_original_and_its_refund(
    ledger: LedgerService, tmp_path: Path, first_kind: str, second_amount: str
) -> None:
    original = uid()
    path = csv_file(
        tmp_path / "conflicting-original.csv",
        ("transaction_id", "kind", "amount_minor", "occurred_on", "original_transaction_id"),
        (
            (original, first_kind, "10000", "2026-10-02", ""),
            (original, "expense", second_amount, "2026-10-02", ""),
            (uid(), "expense_refund", "1000", "2026-10-02", original),
        ),
    )
    service = ExchangeService(ledger)
    before = counts(ledger)
    preview = service.preview(read_table(path), mapping(ledger))
    first, second, refund = preview.rows
    assert all("IMPORT_SOURCE_ID_CONFLICT" in row.issues for row in (first, second))
    assert all(row.fields is None for row in (first, second))
    assert "ORIGINAL_EXPENSE_REQUIRED" in refund.issues and refund.fields is None
    for selection in ([2, 4], [3, 4], [2, 3, 4]):
        with pytest.raises(LedgerError, match="IMPORT_SELECTION_INVALID"):
            service.commit_payload(preview, selection)
    assert counts(ledger) == before and ledger.total_assets() == 100_000


def test_same_source_id_with_different_external_identity_blocks_both_rows(
    ledger: LedgerService, tmp_path: Path
) -> None:
    source_id = uid()
    path = csv_file(
        tmp_path / "conflicting-external-identity.csv",
        (
            "transaction_id",
            "kind",
            "amount_minor",
            "occurred_on",
            "external_source",
            "external_transaction_id",
        ),
        (
            (source_id, "expense", "10000", "2026-10-02", "bank:account-a", "order-one"),
            (source_id, "expense", "10000", "2026-10-02", "bank:account-b", "order-two"),
        ),
    )
    service = ExchangeService(ledger)
    before = counts(ledger)
    preview = service.preview(read_table(path), mapping(ledger))
    assert len(preview.rows) == 2
    assert all(row.fields is None for row in preview.rows)
    assert all("IMPORT_SOURCE_ID_CONFLICT" in row.issues for row in preview.rows)
    for selected in ([2], [3], [2, 3]):
        with pytest.raises(LedgerError, match="IMPORT_SELECTION_INVALID"):
            service.commit_payload(preview, selected)
    assert counts(ledger) == before and ledger.total_assets() == 100_000


@pytest.mark.parametrize("explicit_external", [False, True])
def test_exact_blocked_original_rebinds_new_refund_to_existing_external_expense(
    ledger: LedgerService, tmp_path: Path, explicit_external: bool
) -> None:
    original = uid()
    external_source = "bank:test-account" if explicit_external else ""
    external_id = "merchant-order-123" if explicit_external else ""
    headers = (
        "transaction_id",
        "kind",
        "amount_minor",
        "occurred_on",
        "original_transaction_id",
        "external_source",
        "external_transaction_id",
    )
    path = csv_file(
        tmp_path / "initial-external.csv",
        headers,
        ((original, "expense", "10000", "2026-10-02", "", external_source, external_id),),
    )
    service = ExchangeService(ledger)
    initial = service.preview(read_table(path), mapping(ledger))
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(initial, [2]))
    existing = initial.rows[0].transaction_id
    assert existing != original and ledger.total_assets() == 90_000
    new_path = csv_file(
        tmp_path / "new-refund.csv",
        headers,
        (
            (original, "expense", "10000", "2026-10-02", "", external_source, external_id),
            (uid(), "expense_refund", "1000", "2026-10-02", original, "", ""),
        ),
    )
    preview = service.preview(read_table(new_path), mapping(ledger))
    duplicate, refund = preview.rows
    assert duplicate.exact_duplicate and "DUPLICATE_IMPORT" in duplicate.issues
    assert not refund.issues
    assert cast(dict[str, object], refund.fields)["original_transaction_id"] == existing
    ledger.execute(uid(), "import.commit.v1", service.commit_payload(preview, [3]))
    assert ledger.total_assets() == 91_000
    assert ledger.transaction(refund.transaction_id)["original_transaction_id"] == existing
    assert len(ledger.import_batches()) == 2


def copied_destination(source: LedgerService, root: Path) -> LedgerService:
    database = Database(root / "new.sqlite3")
    database.initialize()
    result = LedgerService(database, clock=source.clock)
    result.ensure_defaults()
    for account in source.entities("account"):
        result.execute(
            uid(),
            "account.create.v1",
            {
                "id": uid(),
                "name": account["name"],
                "account_type": account["account_type"],
                "balance_start_on": account["balance_start_on"],
                "opening_balance_minor": 0,
            },
        )
    for tag in source.entities("tag"):
        result.execute(uid(), "tag.create.v1", {"id": uid(), "name": tag["name"]})
    return result


def record_four_kinds(ledger: LedgerService) -> tuple[str, ...]:
    r = refs(ledger)
    tag = uid()
    ledger.execute(uid(), "tag.create.v1", {"id": tag, "name": "朋友,出游"})
    expense = record(
        ledger,
        amount=1_000,
        occurred_on=date(2026, 9, 30),
        time_zone="Asia/Shanghai",
        occurrence_precision="exact",
        occurred_at_utc=datetime(2026, 9, 30, 4, 30, 0, 123000, tzinfo=UTC),
        note="=SUM(1,2)\n'原样",
        source="local_rule",
        source_text="昨天下午咖啡10元，和朋友",
        tag_ids=(tag,),
        counterparty="朋友",
        merchant="咖啡店",
        location="上海",
    )
    income = record(ledger, amount=5_000, kind="income", note="+收入", tag_ids=(tag,))
    transfer, refund = uid(), uid()
    ledger.transfer(
        TransferFields(
            r["cash"],
            r["bank"],
            2_000,
            TODAY,
            occurrence_precision="period",
            time_period="afternoon",
            note="@转账",
            tag_ids=(tag,),
        ),
        request_id=uid(),
        transaction_id=transfer,
    )
    ledger.refund(
        RefundFields(
            expense, 100, r["cash"], TODAY, note="-退款", tag_ids=(tag,), source_text="退回1元"
        ),
        request_id=uid(),
        transaction_id=refund,
    )
    return expense, income, transfer, refund


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_four_kind_roundtrip_maps_new_account_names_and_preserves_precision(
    ledger: LedgerService, tmp_path: Path, format: str
) -> None:
    identifiers = record_four_kinds(ledger)
    path = tmp_path / ("roundtrip." + format)
    exported = ExchangeService(ledger).export_transactions(path, format)
    assert exported.row_count == 5  # Four financial rows plus the source opening baseline.
    destination = copied_destination(ledger, tmp_path)
    assert refs(destination)["cash"] != refs(ledger)["cash"]
    service = ExchangeService(destination)
    preview = service.preview(read_table(path), ExchangeMapping())
    ordinary = [row for row in preview.rows if row.fields is not None and not row.issues]
    assert len(ordinary) == 4
    assert sum("UNSUPPORTED_IMPORT_KIND" in row.issues for row in preview.rows) == 1
    ledger_payload = service.commit_payload(preview, [row.source_row_number for row in ordinary])
    destination.execute(uid(), "import.commit.v1", ledger_payload)
    assert destination.total_assets() == 4_100  # A portable import does not replay the opening.
    for candidate in ordinary:
        assert candidate.external_transaction_id in identifiers
        source = ledger.transaction(str(candidate.external_transaction_id))
        imported = destination.transaction(candidate.transaction_id)
        for key in (
            "kind",
            "amount_minor",
            "occurred_on",
            "occurrence_precision",
            "time_period",
            "occurred_at_utc",
            "time_zone",
            "note",
            "source_text",
        ):
            assert imported[key] == source[key]
        assert imported["source"] == "import" and imported["tag_ids"]
    # Re-exporting the source into its own database must flag every supported row exactly.
    own = ExchangeService(ledger).preview(read_table(path), ExchangeMapping())
    assert sum(row.exact_duplicate for row in own.rows) == 4


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_formula_injection_is_exported_as_literal_and_reversible(
    ledger: LedgerService, tmp_path: Path, format: str
) -> None:
    text = '=HYPERLINK("https://example.invalid","x")\n\'原始'
    record(ledger, note=text, source_text="@不是公式")
    path = tmp_path / ("literal." + format)
    ExchangeService(ledger).export_transactions(path, format, TransactionFilter(kind="expense"))
    if format == "csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            raw = next(csv.DictReader(handle))
        assert raw["note"] == "'" + text
        assert raw["text_escape"] == "apostrophe-v1"
    else:
        workbook = load_workbook(path, data_only=False)
        try:
            sheet = cast(Worksheet, workbook.active)
            headers = [cell.value for cell in sheet[1]]
            cell = sheet.cell(2, headers.index("note") + 1)
            assert cell.value == text and cell.data_type == "s"
        finally:
            workbook.close()
    destination = copied_destination(ledger, tmp_path)
    preview = ExchangeService(destination).preview(read_table(path), ExchangeMapping())
    assert not preview.rows[0].issues
    assert cast(dict[str, object], preview.rows[0].fields)["note"] == text
    assert cast(dict[str, object], preview.rows[0].fields)["source_text"] == "@不是公式"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_cancellation_and_io_failure_preserve_existing_file(
    ledger: LedgerService, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, format: str
) -> None:
    record(ledger)
    path = tmp_path / ("existing." + format)
    original = b"original unchanged bytes"
    path.write_bytes(original)
    service = ExchangeService(ledger)
    calls = 0

    def cancel_late() -> bool:
        nonlocal calls
        calls += 1
        return calls >= 4

    with pytest.raises(LedgerError, match="EXCHANGE_CANCELLED"):
        service.export_transactions(path, format, cancel=cancel_late)
    assert path.read_bytes() == original
    assert not tuple(tmp_path.glob(".openledger-*"))

    def fail_replace(_source: object, _destination: object) -> None:
        raise PermissionError("injected publication failure")

    monkeypatch.setattr("openledger.infrastructure.exchange.os.replace", fail_replace)
    with pytest.raises(LedgerError, match="EXPORT_IO_ERROR"):
        service.export_transactions(path, format)
    assert path.read_bytes() == original
    assert not tuple(tmp_path.glob(".openledger-*"))


def test_export_is_one_snapshot_even_when_writer_commits_after_read(
    ledger: LedgerService, tmp_path: Path
) -> None:
    first = record(ledger)
    calls = 0
    late = ""
    with ledger.database.read() as connection:
        before_revision = int(connection.execute("SELECT MAX(seq) FROM change_log").fetchone()[0])

    def change_after_snapshot() -> bool:
        nonlocal calls, late
        calls += 1
        if calls == 2:
            late = record(ledger, amount=700)
        return False

    path = tmp_path / "snapshot.csv"
    result = ExchangeService(ledger).export_transactions(
        path, "csv", TransactionFilter(kind="expense"), cancel=change_after_snapshot
    )
    assert result.row_count == 1 and result.data_revision == before_revision
    assert late and late != first
    destination = copied_destination(ledger, tmp_path)
    preview = ExchangeService(destination).preview(read_table(path), ExchangeMapping())
    assert {row.external_transaction_id for row in preview.rows} == {first}


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_large_integer_amounts_stay_decimal_strings(
    ledger: LedgerService, tmp_path: Path, format: str
) -> None:
    identifier = record(ledger, amount=MAX_EVENT_MINOR, kind="income")
    path = tmp_path / ("large." + format)
    ExchangeService(ledger).export_transactions(path, format, TransactionFilter(kind="income"))
    destination = copied_destination(ledger, tmp_path)
    preview = ExchangeService(destination).preview(read_table(path), ExchangeMapping())
    assert not preview.rows[0].issues
    assert preview.rows[0].external_transaction_id == identifier
    assert cast(dict[str, object], preview.rows[0].fields)["amount_minor"] == MAX_EVENT_MINOR


def test_opening_and_adjustment_export_are_inspection_only(
    ledger: LedgerService, tmp_path: Path
) -> None:
    r = refs(ledger)
    ledger.execute(
        uid(),
        "account.adjust.v1",
        {
            "account_id": r["cash"],
            "target_balance_minor": 101_000,
            "occurred_on": TODAY,
            "time_zone": "UTC",
            "reason": "测试校准",
        },
    )
    path = tmp_path / "baseline.csv"
    ExchangeService(ledger).export_transactions(path, "csv")
    destination = copied_destination(ledger, tmp_path)
    preview = ExchangeService(destination).preview(read_table(path), ExchangeMapping())
    assert len(preview.rows) == 2
    assert all("UNSUPPORTED_IMPORT_KIND" in row.issues for row in preview.rows)
    assert all(row.fields is None for row in preview.rows)


def test_error_report_contains_only_selected_diagnostics(
    ledger: LedgerService, tmp_path: Path
) -> None:
    preview = preview_csv(
        ledger, tmp_path / "error.csv", (("支出", "bad", "2026-10-02", "私密原文"),)
    )
    output = tmp_path / "errors.csv"
    ExchangeService(ledger).export_errors(preview, output)
    with output.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["source_row"] == "2" and "INVALID_AMOUNT" in rows[0]["issues"]
    assert "私密原文" not in output.read_text(encoding="utf-8-sig")
