"""Actual desktop file review, same-writer commits, rollback and report drill-down."""

import csv
from collections.abc import Callable, Iterator
from concurrent.futures import Future
from dataclasses import replace
from datetime import date
from pathlib import Path
from threading import Event
from time import monotonic
from typing import cast
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import QFileDialog, QMessageBox, QTableWidget
from pytestqt.qtbot import QtBot

from openledger.application.dto.exchange import ExchangeMapping, FileTable, ImportPreview
from openledger.application.dto.ledger import RefundFields, TransactionFields
from openledger.application.dto.queries import TransactionFilter
from openledger.application.dto.runtime import RuntimeInfo
from openledger.domain.errors import LedgerError
from openledger.infrastructure.ledger import LedgerService
from openledger.infrastructure.settings import SettingsStore
from openledger.presentation.views import exchange as exchange_view
from openledger.presentation.views.main_window import MainWindow

pytestmark = pytest.mark.ui
TODAY = date(2026, 10, 2)


def uid() -> str:
    return str(uuid4())


def refs(ledger: LedgerService) -> dict[str, str]:
    return {
        "cash": str(
            next(row["id"] for row in ledger.entities("account") if row["account_type"] == "cash")
        ),
        "bank": str(
            next(row["id"] for row in ledger.entities("account") if row["account_type"] == "bank")
        ),
        "book": str(ledger.entities("book")[0]["id"]),
        "expense": str(
            next(row["id"] for row in ledger.entities("category") if row["name"] == "餐饮")
        ),
        "income": str(
            next(row["id"] for row in ledger.entities("category") if row["name"] == "工资")
        ),
    }


def item(
    ledger: LedgerService, *, amount: int = 2500, note: str | None = None
) -> TransactionFields:
    values = refs(ledger)
    return TransactionFields(
        "expense", amount, values["cash"], values["book"], values["expense"], TODAY, note=note
    )


def record(ledger: LedgerService, fields: TransactionFields) -> str:
    identifier = uid()
    ledger.record(fields, request_id=uid(), transaction_id=identifier)
    return identifier


def source_file(
    path: Path, rows: tuple[tuple[str, ...], ...] = (("支出", "25", "2026-10-02", "咖啡"),)
) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("类型", "金额", "日期", "备注"))
        writer.writerows(rows)
    return path


def counts(ledger: LedgerService) -> tuple[int, ...]:
    with ledger.database.read() as connection:
        return tuple(
            int(connection.execute("SELECT COUNT(*) FROM " + table).fetchone()[0])
            for table in (
                "transactions",
                "account_entries",
                "import_batches",
                "audit_events",
                "change_log",
                "command_receipts",
            )
        )


def cell(table: QTableWidget, row: int, column: int) -> str:
    value = table.item(row, column)
    assert value is not None
    return value.text()


def finish(qtbot: QtBot, window: MainWindow) -> None:
    qtbot.waitUntil(lambda: not window.exchange.tasks.busy and not window.bridge.busy, timeout=5000)


def defaults(window: MainWindow, ledger: LedgerService) -> None:
    values = refs(ledger)
    for key, identifier in (
        ("book_id", values["book"]),
        ("account_id", values["cash"]),
        ("expense_category_id", values["expense"]),
        ("income_category_id", values["income"]),
    ):
        combo = window.exchange.targets[key]
        combo.setCurrentIndex(combo.findData(identifier))


def preview(qtbot: QtBot, window: MainWindow, ledger: LedgerService, path: Path) -> ImportPreview:
    window.exchange.load_file(path)
    finish(qtbot, window)
    defaults(window, ledger)
    window.exchange.build_preview()
    finish(qtbot, window)
    value = window.exchange.preview
    assert value is not None, window.exchange.status.text()
    return value


@pytest.fixture
def window(qtbot: QtBot, ledger: LedgerService, tmp_path: Path) -> Iterator[MainWindow]:
    runtime = RuntimeInfo(
        "0.2.0.dev0", "3.12.5", "6.11.2", "6.11.2", "3.45.3", str(tmp_path), False
    )
    widget = MainWindow(runtime, ledger, SettingsStore(tmp_path / "settings.json"))
    qtbot.addWidget(widget)
    widget.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
    widget.show()
    yield widget
    widget.close()


def test_choose_headers_default_mapping_preview_and_same_writer_import_revert(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = source_file(
        tmp_path / "中文账单.csv",
        (("支出", "25", "2026-10-02", "咖啡"), ("收入", "50", "2026-10-02", "工资")),
    )
    before_counts, before_assets = counts(ledger), ledger.total_assets()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args, **kwargs: (str(path), ""))
    window.exchange._choose_file()
    finish(qtbot, window)
    assert window.exchange.source_table is not None
    assert window.exchange.columns["amount"].currentData() == "金额"
    assert window.exchange.columns["kind"].currentData() == "类型"
    assert counts(ledger) == before_counts and ledger.total_assets() == before_assets
    defaults(window, ledger)
    window.exchange.build_preview()
    finish(qtbot, window)
    assert window.exchange.preview is not None and len(window.exchange.preview.rows) == 2
    assert counts(ledger) == before_counts
    for number in range(2):
        selected = window.exchange.rows.item(number, 0)
        assert selected is not None and selected.checkState() == Qt.CheckState.Checked
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes
    )
    window.exchange.confirm_import()
    finish(qtbot, window)
    assert "已保存" in window.notice.text()
    assert window.exchange.preview is None
    assert ledger.total_assets() == before_assets + 2500
    batches = ledger.import_batches()
    assert len(batches) == 1 and batches[0]["accepted_row_count"] == 2
    assert cell(window.exchange.batches, 0, 0) == path.name
    assert cell(window.exchange.batches, 0, 2) == "已提交"
    assert window.queries.transactions(TransactionFilter(kind="expense")).total == 1
    assert window.queries.transactions(TransactionFilter(kind="income")).total == 1
    window.exchange.batches.selectRow(0)
    window.exchange.revert_batch()
    finish(qtbot, window)
    assert ledger.total_assets() == before_assets
    assert ledger.import_batches()[0]["status"] == "reverted"
    assert cell(window.exchange.batches, 0, 2) == "已撤销"
    assert window.queries.transactions(TransactionFilter(kind="expense")).total == 0


def test_invalid_row_disabled_possible_duplicate_unchecked_and_fresh_row_checked(
    qtbot: QtBot, window: MainWindow, ledger: LedgerService, tmp_path: Path
) -> None:
    record(ledger, item(ledger))
    plan = preview(
        qtbot,
        window,
        ledger,
        source_file(
            tmp_path / "检查.csv",
            (
                ("支出", "25", "2026-10-02", "疑似"),
                ("支出", "1.234", "2026-10-02", "错误"),
                ("支出", "30", "2026-10-02", "新记录"),
            ),
        ),
    )
    states = [window.exchange.rows.item(number, 0) for number in range(3)]
    assert all(state is not None for state in states)
    fuzzy, invalid, valid = states
    assert fuzzy is not None and invalid is not None and valid is not None
    assert (
        fuzzy.checkState() == Qt.CheckState.Unchecked and fuzzy.flags() & Qt.ItemFlag.ItemIsEnabled
    )
    assert (
        invalid.checkState() == Qt.CheckState.Unchecked
        and invalid.flags() == Qt.ItemFlag.NoItemFlags
    )
    assert valid.checkState() == Qt.CheckState.Checked
    assert plan.rows[0].possible_duplicates and plan.rows[1].issues == ("AMOUNT_PRECISION",)
    assert "可能重复" in cell(window.exchange.rows, 0, window.exchange.rows.columnCount() - 1)


def test_default_targets_are_explicit_and_missing_fields_cannot_be_selected(
    qtbot: QtBot, window: MainWindow, tmp_path: Path
) -> None:
    window.exchange.load_file(source_file(tmp_path / "缺少目标.csv"))
    finish(qtbot, window)
    assert all(combo.currentData() is None for combo in window.exchange.targets.values())
    window.exchange.build_preview()
    finish(qtbot, window)
    assert window.exchange.preview is not None and window.exchange.preview.rows[0].issues
    selected = window.exchange.rows.item(0, 0)
    assert selected is not None and selected.flags() == Qt.ItemFlag.NoItemFlags
    window.exchange.confirm_import()
    assert "至少一行" in window.exchange.status.text()


def test_confirmation_decline_does_not_mutate_and_keeps_preview(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = preview(qtbot, window, ledger, source_file(tmp_path / "暂不导入.csv"))
    before = counts(ledger)
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.No
    )
    window.exchange.confirm_import()
    assert window.exchange.preview is plan and counts(ledger) == before and not window.bridge.busy


def test_cancel_after_payload_completes_before_gui_delivery_blocks_ledger_command(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = preview(qtbot, window, ledger, source_file(tmp_path / "取消待送达批次.csv"))
    original_finished = window.exchange.tasks._finished
    queued = Event()
    commands: list[tuple[str, object]] = []
    window.exchange.commandRequested.connect(
        lambda command, payload: commands.append((command, payload))
    )
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes
    )

    def observed_finished(future: Future[object]) -> None:
        original_finished(future)
        queued.set()

    monkeypatch.setattr(window.exchange.tasks, "_finished", observed_finished)
    before, assets = counts(ledger), ledger.total_assets()
    window.exchange.confirm_import()
    # Waiting on the threading event deliberately does not process queued Qt delivery.
    assert queued.wait(3) and window.exchange.tasks.busy
    assert window.exchange._preparing_commit
    window.exchange.cancel_task()
    finish(qtbot, window)
    assert commands == [] and not window.bridge.busy
    assert counts(ledger) == before and ledger.total_assets() == assets
    assert window.exchange.preview is plan and window.exchange.commit_button.isEnabled()


def test_failed_commit_preserves_review_and_successful_retry_uses_same_plan(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = preview(qtbot, window, ledger, source_file(tmp_path / "重试.csv"))
    before = counts(ledger)
    original = ledger._fault

    def fail(stage: str) -> None:
        if stage == "before_commit":
            raise LedgerError("DATABASE_BUSY")

    ledger._fault = fail
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes
    )
    try:
        window.exchange.confirm_import()
        finish(qtbot, window)
        assert window.exchange.preview is plan
        assert (
            "预览已保留" in window.exchange.status.text()
            and "DATABASE_BUSY" in window.notice.text()
        )
        assert counts(ledger) == before and ledger.import_batches() == ()
    finally:
        ledger._fault = original
    window.exchange.confirm_import()
    finish(qtbot, window)
    assert ledger.import_batches()[0]["id"] == plan.batch_id and "已保存" in window.notice.text()


def test_mapping_change_invalidates_existing_confirmation(
    qtbot: QtBot, window: MainWindow, ledger: LedgerService, tmp_path: Path
) -> None:
    preview(qtbot, window, ledger, source_file(tmp_path / "映射改变.csv"))
    window.exchange.columns["note"].setCurrentIndex(0)
    assert window.exchange.preview is None and window.exchange.rows.rowCount() == 0
    assert not window.exchange.commit_button.isEnabled()


def test_outdated_preview_generation_is_discarded(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    window.exchange.load_file(source_file(tmp_path / "过期预览.csv"))
    finish(qtbot, window)
    defaults(window, ledger)
    original = window.exchange.service.preview
    entered, release = Event(), Event()

    def delayed(
        table: FileTable, mapping: ExchangeMapping, *, cancel: Callable[[], bool] | None = None
    ) -> ImportPreview:
        entered.set()
        assert release.wait(3)
        return original(table, mapping)

    monkeypatch.setattr(window.exchange.service, "preview", delayed)
    window.exchange.build_preview()
    qtbot.waitUntil(entered.is_set)
    try:
        window.exchange.columns["note"].setCurrentIndex(0)
    finally:
        release.set()
    finish(qtbot, window)
    assert window.exchange.preview is None and window.exchange.rows.rowCount() == 0
    assert not window.exchange.commit_button.isEnabled()


def test_exact_duplicate_is_blocked_after_previous_import(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = source_file(tmp_path / "同一文件.csv")
    preview(qtbot, window, ledger, path)
    monkeypatch.setattr(
        QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes
    )
    window.exchange.confirm_import()
    finish(qtbot, window)
    plan = preview(qtbot, window, ledger, path)
    assert plan.rows[0].exact_duplicate and "DUPLICATE_IMPORT" in plan.rows[0].issues
    selected = window.exchange.rows.item(0, 0)
    assert selected is not None and selected.flags() == Qt.ItemFlag.NoItemFlags


@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
def test_native_export_dialog_writes_real_snapshot_file(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    format_name: str,
) -> None:
    record(ledger, item(ledger, note="=SUM(A1:A2)"))
    target = tmp_path / ("中文导出." + format_name)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(target), ""))
    window.exchange._export_transactions(format_name)
    finish(qtbot, window)
    assert "已导出 2 行" in window.exchange.status.text() and target.stat().st_size > 100
    if format_name == "csv":
        with target.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        expense = next(row for row in rows if row["kind"] == "'expense")
        assert expense["amount_minor"] == "'2500" and expense["note"] == "'=SUM(A1:A2)"
    else:
        workbook = load_workbook(target, read_only=True, data_only=False)
        try:
            sheet = workbook.active
            assert sheet is not None
            worksheet_rows = list(sheet.iter_rows(values_only=True))
            note_column = worksheet_rows[0].index("note")
            assert any(row[note_column] == "=SUM(A1:A2)" for row in worksheet_rows[1:])
            for row in sheet.iter_rows(min_row=2):
                assert all(cell.data_type != "f" for cell in row)
        finally:
            workbook.close()


@pytest.mark.parametrize("format_name", ["csv", "xlsx"])
def test_failed_atomic_export_preserves_previous_file(
    qtbot: QtBot,
    window: MainWindow,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    format_name: str,
) -> None:
    target = tmp_path / ("原文件." + format_name)
    target.write_bytes(b"previous exact bytes")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(target), ""))

    def fail_replace(source: object, destination: object) -> None:
        raise OSError("simulated publication failure")

    monkeypatch.setattr("openledger.infrastructure.exchange.os.replace", fail_replace)
    window.exchange._export_transactions(format_name)
    finish(qtbot, window)
    assert target.read_bytes() == b"previous exact bytes"
    assert "EXPORT_IO_ERROR" in window.exchange.status.text()
    assert not tuple(tmp_path.glob(".openledger-*"))


def test_error_details_export_contains_invalid_row_codes(
    qtbot: QtBot,
    window: MainWindow,
    ledger: LedgerService,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    preview(
        qtbot,
        window,
        ledger,
        source_file(tmp_path / "错误行.csv", (("支出", "1.234", "2026-10-02", "错误"),)),
    )
    target = tmp_path / "错误明细.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args, **kwargs: (str(target), ""))
    window.exchange._export_errors()
    finish(qtbot, window)
    with target.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[0]["source_row"] == "2" and rows[0]["issues"] == "AMOUNT_PRECISION"


def test_cancel_file_read_and_close_workers_do_not_publish_stale_table(
    qtbot: QtBot, window: MainWindow, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    entered, stopped = Event(), Event()

    def cancellable(
        path: Path, mapping: ExchangeMapping, *, cancel: Callable[[], bool] | None = None
    ) -> FileTable:
        entered.set()
        deadline = monotonic() + 3
        while monotonic() < deadline:
            if cancel is not None and cancel():
                stopped.set()
                raise LedgerError("EXCHANGE_CANCELLED")
            Event().wait(0.01)
        pytest.fail("file operation did not stop")

    monkeypatch.setattr(exchange_view, "read_table", cancellable)
    window.exchange.load_file(tmp_path / "cancel.csv")
    qtbot.waitUntil(entered.is_set)
    window.exchange.cancel_task()
    finish(qtbot, window)
    assert stopped.is_set() and window.exchange.source_table is None
    assert "EXCHANGE_CANCELLED" in window.exchange.status.text()
    entered.clear()
    stopped.clear()
    window.exchange.load_file(tmp_path / "close.csv")
    qtbot.waitUntil(entered.is_set)
    window.exchange.close_workers()
    assert stopped.is_set() and not window.exchange.tasks.busy
    assert window.exchange.source_table is None


def test_category_drilldown_retains_complete_report_scope_and_refund_net(
    qtbot: QtBot, window: MainWindow, ledger: LedgerService
) -> None:
    values = refs(ledger)
    book, category, tag, other_tag = uid(), uid(), uid(), uid()
    ledger.execute(uid(), "book.create.v1", {"id": book, "name": "旅行"})
    ledger.execute(uid(), "category.create.v1", {"id": category, "name": "车费", "kind": "expense"})
    for identifier, name in ((tag, "旅行"), (other_tag, "其他")):
        ledger.execute(uid(), "tag.create.v1", {"id": identifier, "name": name})
    base = replace(item(ledger, amount=500), book_id=book, category_id=category, tag_ids=(tag,))
    original = record(ledger, base)
    record(ledger, replace(base, amount_minor=700, occurred_on=date(2026, 10, 1)))
    record(ledger, replace(base, account_id=values["bank"]))
    record(ledger, replace(base, tag_ids=(other_tag,)))
    ledger.refund(
        RefundFields(original, 200, values["cash"], TODAY, tag_ids=(tag,)),
        request_id=uid(),
        transaction_id=uid(),
    )
    window.refresh()
    qtbot.waitUntil(
        lambda: not window.analysis._reports.busy and window.analysis._pending is None, timeout=5000
    )
    window.analysis.start.setDate(QDate(2026, 10, 2))
    for combo, identifier in (
        (window.analysis.book, book),
        (window.analysis.account, values["cash"]),
        (window.analysis.category, category),
        (window.analysis.tag, tag),
    ):
        combo.setCurrentIndex(combo.findData(identifier))
    window.analysis.refresh()
    qtbot.waitUntil(
        lambda: not window.analysis._reports.busy and window.analysis._pending is None, timeout=5000
    )
    report = window.analysis.report
    assert report is not None and report.totals.net_expense_minor == 300
    window.transactions.search.setText("stale unrelated filter")
    window.transactions.include_deleted.setChecked(True)
    window.analysis.categoryRequested.emit(category)
    transactions = window.transactions
    assert transactions.start.date().toPython() == transactions.end.date().toPython() == TODAY
    assert (
        transactions.book.currentData() == book
        and transactions.account.currentData() == values["cash"]
    )
    assert transactions.category.currentData() == category and transactions.tag.currentData() == tag
    assert transactions.search.text() == "" and not transactions.include_deleted.isChecked()
    rows = transactions._rows
    assert len(rows) == 2
    net = sum(
        cast(int, row["amount_minor"]) * (-1 if row["kind"] == "expense_refund" else 1)
        for row in rows
    )
    assert net == report.totals.net_expense_minor
    assert window._pages.currentWidget() is transactions
