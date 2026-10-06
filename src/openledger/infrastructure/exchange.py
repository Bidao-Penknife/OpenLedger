"""Bounded file readers, deterministic import previews and atomic snapshot exports."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import tempfile
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import asdict, replace
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import cast
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from openledger.application.dto.exchange import (
    ExchangeMapping,
    ExportResult,
    FileTable,
    ImportPreview,
    ImportRow,
)
from openledger.application.dto.queries import TransactionFilter
from openledger.domain.currencies import currency
from openledger.domain.errors import LedgerError
from openledger.domain.money import parse_amount, validate_minor
from openledger.domain.values import validate_occurrence
from openledger.infrastructure.ledger import LedgerService
from openledger.infrastructure.queries import (
    _EFFECTIVE_BOOK,
    _EFFECTIVE_CATEGORY,
    _JOINS,
    _change_seq,
    _decorate,
    _validate,
    _where,
)

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 80
MAX_CELL = 8192
Cancel = Callable[[], bool] | None
_DEFAULT_MAPPING = ExchangeMapping()
_DEFAULT_FILTER = TransactionFilter()
EXCHANGE_COLUMNS = (
    "format_version",
    "text_escape",
    "transaction_id",
    "kind",
    "currency_code",
    "amount_minor",
    "to_amount_minor",
    "to_currency_code",
    "occurred_on",
    "occurrence_precision",
    "time_period",
    "occurred_at_utc",
    "time_zone",
    "book_id",
    "book_name",
    "category_id",
    "category_name",
    "account_id",
    "account_name",
    "from_account_id",
    "from_account_name",
    "to_account_id",
    "to_account_name",
    "payment_method_id",
    "payment_method_name",
    "original_transaction_id",
    "tag_ids",
    "tag_names",
    "counterparty",
    "merchant",
    "location",
    "note",
    "source",
    "source_text",
    "external_source",
    "external_transaction_id",
    "balance_before_minor",
    "balance_target_minor",
    "adjustment_reason",
    "deleted_at_utc",
    "version",
)
ALIASES = {
    "类型": "kind",
    "收支类型": "kind",
    "金额": "amount",
    "币种": "currency_code",
    "转入金额": "to_amount",
    "转入币种": "to_currency_code",
    "金额(元)": "amount",
    "日期": "occurred_on",
    "交易日期": "occurred_on",
    "备注": "note",
    "商户": "merchant",
    "对象": "counterparty",
    "地点": "location",
    "账户": "account_name",
    "账本": "book_name",
    "分类": "category_name",
    "支付方式": "payment_method_name",
    "订单号": "external_transaction_id",
    "转出账户": "from_account_name",
    "转入账户": "to_account_name",
}
_KIND_NAMES = {"收入": "income", "支出": "expense", "转账": "transfer", "退款": "expense_refund"}
_TABLES = {
    "book": "books",
    "account": "accounts",
    "category": "categories",
    "tag": "tags",
    "payment_method": "payment_methods",
}


def _cancelled(cancel: Cancel) -> None:
    if cancel is not None and cancel():
        raise LedgerError("EXCHANGE_CANCELLED")


def _source_bytes(path: Path) -> bytes:
    if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
        raise LedgerError("IMPORT_FILE_TOO_LARGE")
    with path.open("rb") as handle:
        data = handle.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise LedgerError("IMPORT_FILE_TOO_LARGE")
    return data


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise LedgerError("IMPORT_INVALID_CELL")
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00":
            return value.date().isoformat()
        return value.isoformat()
    result = str(value)
    if len(result) > MAX_CELL or "\x00" in result:
        raise LedgerError("IMPORT_INVALID_CELL")
    return result


def _check_xlsx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            files = archive.infolist()
            names: set[str] = set()
            if len(files) > 2000 or sum(item.file_size for item in files) > 100 * 1024 * 1024:
                raise LedgerError("IMPORT_ARCHIVE_UNSAFE")
            for item in files:
                name = item.orig_filename
                path = PurePosixPath(name)
                if (
                    name in names
                    or path.is_absolute()
                    or ".." in path.parts
                    or "\\" in name
                    or ":" in name
                    or item.flag_bits & 1
                    or item.file_size > max(item.compress_size, 1) * 200
                ):
                    raise LedgerError("IMPORT_ARCHIVE_UNSAFE")
                names.add(name)
            if archive.testzip() is not None:
                raise LedgerError("IMPORT_ARCHIVE_UNSAFE")
    except zipfile.BadZipFile as error:
        raise LedgerError("IMPORT_INVALID_FILE") from error


def read_table(
    path: Path, mapping: ExchangeMapping = _DEFAULT_MAPPING, *, cancel: Cancel = None
) -> FileTable:
    """Read CSV/xlsx without evaluating formulas or resolving external links."""
    _cancelled(cancel)
    data = _source_bytes(path)
    suffix = path.suffix.lower()
    rows: list[tuple[str, ...]] = []
    formula_rows: set[int] = set()
    sheets: tuple[str, ...] = ()
    try:
        if suffix == ".csv":
            if mapping.encoding not in {"utf-8-sig", "utf-8", "gb18030"}:
                raise LedgerError("IMPORT_ENCODING_INVALID")
            text = data.decode(mapping.encoding)
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
                delimiter = dialect.delimiter
            except csv.Error:
                first_line = text.splitlines()[0] if text else ""
                delimiter = max(
                    (",", ";", "\t"),
                    key=lambda value: len(next(csv.reader([first_line], delimiter=value))),
                )
            csv.field_size_limit(MAX_CELL)
            for item in csv.reader(
                io.StringIO(text, newline=""),
                delimiter=delimiter,
                quotechar='"',
                doublequote=True,
                strict=True,
            ):
                _cancelled(cancel)
                rows.append(tuple(_cell(value) for value in item))
                if len(rows) > MAX_ROWS + 1 or len(item) > MAX_COLUMNS:
                    raise LedgerError("IMPORT_FILE_TOO_LARGE")
        elif suffix == ".xlsx":
            _check_xlsx(data)
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheets = tuple(workbook.sheetnames)
                sheet = workbook[mapping.sheet] if mapping.sheet else workbook.active
                if sheet is None:
                    raise LedgerError("IMPORT_INVALID_FILE")
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise LedgerError("IMPORT_FILE_TOO_LARGE")
                for number, cells in enumerate(sheet.iter_rows(), 1):
                    _cancelled(cancel)
                    if number > MAX_ROWS + 1:
                        raise LedgerError("IMPORT_FILE_TOO_LARGE")
                    rows.append(tuple(_cell(cell.value) for cell in cells))
                    if any(cell.data_type == "f" for cell in cells):
                        formula_rows.add(number)
            finally:
                workbook.close()
        else:
            raise LedgerError("IMPORT_FORMAT_UNSUPPORTED")
    except (UnicodeError, csv.Error, ValueError, KeyError, OSError) as error:
        raise LedgerError("IMPORT_INVALID_FILE") from error
    if not rows or not rows[0] or 1 in formula_rows:
        raise LedgerError("IMPORT_HEADER_INVALID")
    headers = tuple(value.strip() for value in rows[0])
    if any(not value for value in headers) or len(set(headers)) != len(headers):
        raise LedgerError("IMPORT_HEADER_INVALID")
    width = len(headers)
    if any(len(row) != width for row in rows[1:]):
        raise LedgerError("IMPORT_COLUMN_COUNT")
    if len(rows) == 1:
        raise LedgerError("IMPORT_EMPTY_FILE")
    return FileTable(
        path.resolve(),
        suffix[1:],
        hashlib.sha256(data).hexdigest(),
        headers,
        tuple(rows[1:]),
        frozenset(formula_rows),
        sheets,
    )


def automatic_columns(headers: Iterable[str]) -> tuple[tuple[str, str], ...]:
    """Recognize canonical and common Chinese headers without guessing values."""
    result: dict[str, str] = {}
    for header in headers:
        target = ALIASES.get(header, header)
        if target in EXCHANGE_COLUMNS or target in {"amount", "to_amount"}:
            result.setdefault(target, header)
    return tuple(sorted(result.items()))


def _mapping_json(mapping: ExchangeMapping) -> str:
    value = asdict(mapping)
    value["columns"] = sorted(mapping.columns)
    value["normalization_version"] = 1
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _plan_digest(preview: ImportPreview) -> str:
    value = {
        "file_digest": preview.table.file_digest,
        "mapping_json": preview.mapping_json,
        "mapping_hash": preview.mapping_hash,
        "batch_id": preview.batch_id,
        "rows": [asdict(row) for row in preview.rows],
        "data_revision": preview.data_revision,
    }
    encoded = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _resolve(
    references: dict[str, tuple[dict[str, object], ...]],
    entity: str,
    data: dict[str, str],
    fallback: str | None = None,
    *,
    prefix: str | None = None,
) -> str:
    key = prefix or entity
    identifier, name = data.get(key + "_id", ""), data.get(key + "_name", "")
    rows = references[entity]
    found = next((row for row in rows if row["id"] == identifier), None) if identifier else None
    if found is not None and name and found["name"] != name:
        raise LedgerError("IMPORT_REFERENCE_CONFLICT")
    if found is None and name:
        matches = [row for row in rows if row["name"] == name]
        if len(matches) == 1:
            found = matches[0]
    if found is None and fallback and not identifier and not name:
        found = next((row for row in rows if row["id"] == fallback), None)
    if found is None:
        raise LedgerError("IMPORT_REFERENCE_REQUIRED")
    if found["is_archived"]:
        raise LedgerError("ENTITY_ARCHIVED")
    return str(found["id"])


def _data(table: FileTable, mapping: ExchangeMapping, row: tuple[str, ...]) -> dict[str, str]:
    raw = dict(zip(table.headers, row, strict=True))
    data = {target: raw[header] for target, header in mapping.columns if header in raw}
    if data.get("format_version") == "1":
        escape = data.get("text_escape")
        if escape == "apostrophe-v1":
            data = {
                key: value[1:] if value.startswith("'") else value for key, value in data.items()
            }
        elif escape != "literal-v1":
            raise LedgerError("IMPORT_ESCAPE_INVALID")
    elif data.get("format_version"):
        raise LedgerError("IMPORT_VERSION_UNSUPPORTED")
    return data


class ExchangeService:
    """Preview one snapshot; all financial publication uses ledger commands."""

    def __init__(self, ledger: LedgerService) -> None:
        self.ledger = ledger
        self.database = ledger.database

    def preview(
        self, table: FileTable, mapping: ExchangeMapping, *, cancel: Cancel = None
    ) -> ImportPreview:
        """Validate rows, block exact duplicates and flag possible duplicates."""
        if not mapping.columns:
            mapping = replace(mapping, columns=automatic_columns(table.headers))
        if len({target for target, _ in mapping.columns}) != len(mapping.columns) or any(
            header not in table.headers for _, header in mapping.columns
        ):
            raise LedgerError("IMPORT_MAPPING_INVALID")
        mapping_json = _mapping_json(mapping)
        if len(mapping_json) > 8192:
            raise LedgerError("IMPORT_MAPPING_INVALID")
        mapping_hash = hashlib.sha256(mapping_json.encode("utf-8")).hexdigest()
        rows: list[ImportRow] = []
        new_ids = tuple(str(uuid4()) for _ in table.rows)
        with self.database.read() as connection:
            revision = _change_seq(connection)
            references = {
                entity: tuple(dict(row) for row in connection.execute("SELECT * FROM " + name))
                for entity, name in _TABLES.items()
            }
            source_ids: dict[str, str] = {}
            decoded: list[dict[str, str]] = []
            decode_errors: list[str | None] = []
            for source_values, identifier in zip(table.rows, new_ids, strict=True):
                try:
                    item = _data(table, mapping, source_values)
                    decoded.append(item)
                    decode_errors.append(None)
                    if item.get("transaction_id"):
                        source_ids.setdefault(item["transaction_id"], identifier)
                except LedgerError as error:
                    decoded.append({})
                    decode_errors.append(error.code)
            seen_external: set[tuple[str, str]] = set()
            seen_signatures: dict[tuple[object, ...], str] = {}
            for number, (data, identifier, decode_error) in enumerate(
                zip(decoded, new_ids, decode_errors, strict=True), 2
            ):
                _cancelled(cancel)
                issues: list[str] = []
                exact = False
                fuzzy: tuple[str, ...] = ()
                fields: dict[str, object] | None = None
                external_source = data.get("external_source") or None
                external_id = data.get("external_transaction_id") or None
                source_id = data.get("transaction_id")
                if source_id and not external_id:
                    external_source, external_id = "openledger-exchange-v1", source_id
                try:
                    if decode_error:
                        raise LedgerError(decode_error)
                    if number in table.formula_rows:
                        raise LedgerError("IMPORT_FORMULA_FORBIDDEN")
                    if data.get("deleted_at_utc"):
                        raise LedgerError("IMPORT_DELETED_ROW")
                    kind = _KIND_NAMES.get(
                        data.get("kind", ""), data.get("kind") or mapping.default_kind
                    )
                    if kind not in {"income", "expense", "transfer", "expense_refund"}:
                        raise LedgerError("UNSUPPORTED_IMPORT_KIND")
                    primary = _resolve(
                        references,
                        "account",
                        data,
                        mapping.from_account_id if kind == "transfer" else mapping.account_id,
                        prefix="from_account" if kind == "transfer" else "account",
                    )
                    native_code = str(
                        next(row for row in references["account"] if row["id"] == primary)[
                            "currency_code"
                        ]
                    )
                    code = currency(data.get("currency_code") or native_code).code
                    if code != native_code:
                        raise LedgerError("CURRENCY_MISMATCH")
                    amount_text = data.get("amount_minor", "")
                    if amount_text:
                        if not amount_text.isascii() or not amount_text.isdigit():
                            raise LedgerError("INVALID_AMOUNT")
                        amount = int(amount_text)
                        validate_minor(amount)
                    else:
                        amount = parse_amount(data.get("amount", ""), code)
                    try:
                        day = date.fromisoformat(data.get("occurred_on", ""))
                    except ValueError as error:
                        raise LedgerError("INVALID_DATE") from error
                    zone = data.get("time_zone") or mapping.time_zone
                    precision = data.get("occurrence_precision") or "date"
                    instant_text = data.get("occurred_at_utc") or None
                    try:
                        instant = (
                            datetime.fromisoformat(instant_text.replace("Z", "+00:00"))
                            if instant_text
                            else None
                        )
                    except ValueError as error:
                        raise LedgerError("INVALID_DATE") from error
                    validate_occurrence(
                        day,
                        zone,
                        precision,
                        data.get("time_period") or None,
                        instant,
                        now=self.ledger.clock(),
                    )
                    fields = {
                        "kind": kind,
                        "amount_minor": amount,
                        "occurred_on": day.isoformat(),
                        "time_zone": zone,
                        "occurrence_precision": precision,
                        "time_period": data.get("time_period") or None,
                        "occurred_at_utc": instant_text,
                        "currency_code": code,
                        "source": "import",
                        "source_text": data.get("source_text") or None,
                    }
                    for key in ("counterparty", "merchant", "location", "note"):
                        fields[key] = data.get(key) or None
                    tag_ids: tuple[str, ...] = ()
                    if data.get("tag_names"):
                        names = json.loads(data["tag_names"])
                        if not isinstance(names, list) or not all(
                            isinstance(name, str) for name in names
                        ):
                            raise LedgerError("IMPORT_TAG_INVALID")
                        tag_ids = tuple(
                            _resolve(references, "tag", {"tag_name": name}) for name in names
                        )
                    elif data.get("tag_ids"):
                        values = json.loads(data["tag_ids"])
                        if not isinstance(values, list) or not all(
                            isinstance(item, str) for item in values
                        ):
                            raise LedgerError("IMPORT_TAG_INVALID")
                        tag_ids = tuple(
                            _resolve(references, "tag", {"tag_id": item}) for item in values
                        )
                    fields["tag_ids"] = tag_ids
                    if data.get("payment_method_id") or data.get("payment_method_name"):
                        fields["payment_method_id"] = _resolve(references, "payment_method", data)
                    if kind in {"income", "expense"}:
                        fields["book_id"] = _resolve(references, "book", data, mapping.book_id)
                        category = _resolve(
                            references,
                            "category",
                            data,
                            mapping.income_category_id
                            if kind == "income"
                            else mapping.expense_category_id,
                        )
                        if (
                            next(row for row in references["category"] if row["id"] == category)[
                                "transaction_kind"
                            ]
                            != kind
                        ):
                            raise LedgerError("CATEGORY_KIND_MISMATCH")
                        fields["category_id"] = category
                        fields["account_id"] = _resolve(
                            references, "account", data, mapping.account_id
                        )
                    elif kind == "transfer":
                        fields["from_account_id"] = _resolve(
                            references,
                            "account",
                            data,
                            mapping.from_account_id,
                            prefix="from_account",
                        )
                        fields["to_account_id"] = _resolve(
                            references, "account", data, mapping.to_account_id, prefix="to_account"
                        )
                        if fields["from_account_id"] == fields["to_account_id"]:
                            raise LedgerError("SAME_ACCOUNT_TRANSFER")
                        target_code = str(
                            next(
                                row
                                for row in references["account"]
                                if row["id"] == fields["to_account_id"]
                            )["currency_code"]
                        )
                        if data.get("to_currency_code") and data["to_currency_code"] != target_code:
                            raise LedgerError("CURRENCY_MISMATCH")
                        incoming = data.get("to_amount_minor", "")
                        if incoming:
                            if not incoming.isascii() or not incoming.isdigit():
                                raise LedgerError("INVALID_AMOUNT")
                            fields["to_amount_minor"] = int(incoming)
                            validate_minor(fields["to_amount_minor"])
                        elif data.get("to_amount"):
                            fields["to_amount_minor"] = parse_amount(data["to_amount"], target_code)
                        elif code != target_code:
                            raise LedgerError("TRANSFER_TARGET_AMOUNT_REQUIRED")
                        for key in ("payment_method_id", "counterparty", "merchant", "location"):
                            fields.pop(key, None)
                    else:
                        fields["account_id"] = _resolve(
                            references, "account", data, mapping.account_id
                        )
                        original = data.get("original_transaction_id", "")
                        existing = connection.execute(
                            "SELECT id FROM transactions WHERE id=? AND kind='expense' "
                            "AND deleted_at_utc IS NULL",
                            (original,),
                        ).fetchone()
                        if existing:
                            fields["original_transaction_id"] = existing[0]
                        elif original in source_ids:
                            fields["original_transaction_id"] = source_ids[original]
                        else:
                            match = connection.execute(
                                "SELECT id FROM transactions "
                                "WHERE external_source='openledger-exchange-v1' "
                                "AND external_transaction_id=? AND kind='expense' "
                                "AND deleted_at_utc IS NULL",
                                (original,),
                            ).fetchone()
                            if not match:
                                raise LedgerError("ORIGINAL_EXPENSE_REQUIRED")
                            fields["original_transaction_id"] = match[0]
                        for key in ("counterparty", "merchant", "location"):
                            fields.pop(key, None)
                    account_ids = [
                        cast(str, fields[key])
                        for key in ("account_id", "from_account_id", "to_account_id")
                        if key in fields
                    ]
                    for account_id in account_ids:
                        account = next(
                            row for row in references["account"] if row["id"] == account_id
                        )
                        if day.isoformat() < str(account["balance_start_on"]):
                            raise LedgerError("BEFORE_ACCOUNT_START")
                    if bool(external_source) != bool(external_id):
                        raise LedgerError("IMPORT_EXTERNAL_ID_INVALID")
                    exact = bool(
                        connection.execute(
                            "SELECT 1 FROM transactions WHERE (import_file_digest=? "
                            "AND import_mapping_hash=? AND import_source_row=?) "
                            "OR (external_source=? AND external_transaction_id=?) OR id=?",
                            (
                                table.file_digest,
                                mapping_hash,
                                number,
                                external_source,
                                external_id,
                                source_id,
                            ),
                        ).fetchone()
                    )
                    if external_source and external_id:
                        pair = (external_source, external_id)
                        if pair in seen_external:
                            exact = True
                        seen_external.add(pair)
                    signature = (
                        kind,
                        amount,
                        day.isoformat(),
                        fields.get("account_id"),
                        fields.get("from_account_id"),
                        fields.get("to_account_id"),
                        fields.get("book_id"),
                        fields.get("category_id"),
                    )
                    candidates = connection.execute(
                        "SELECT t.id FROM transactions t WHERE t.deleted_at_utc IS NULL "
                        "AND t.kind=? AND t.amount_minor=? AND t.occurred_on=? "
                        "AND EXISTS(SELECT 1 FROM account_entries e "
                        "WHERE e.transaction_id=t.id AND e.account_id=?) LIMIT 20",
                        (kind, amount, day.isoformat(), account_ids[0]),
                    ).fetchall()
                    fuzzy = tuple(str(row[0]) for row in candidates)
                    if signature in seen_signatures:
                        fuzzy += (seen_signatures[signature],)
                    seen_signatures[signature] = identifier
                    if exact:
                        issues.append("DUPLICATE_IMPORT")
                except (LedgerError, ValueError, TypeError) as error:
                    issues.append(
                        error.code if isinstance(error, LedgerError) else "IMPORT_INVALID_CELL"
                    )
                    fields = None
                rows.append(
                    ImportRow(
                        number,
                        identifier,
                        fields,
                        external_source,
                        external_id,
                        tuple(issues),
                        exact,
                        fuzzy,
                    )
                )
            groups: dict[str, list[int]] = {}
            for index, item in enumerate(decoded):
                if source_id := item.get("transaction_id"):
                    groups.setdefault(source_id, []).append(index)
            for indices in groups.values():
                semantics = {
                    json.dumps(
                        {
                            "fields": rows[index].fields,
                            "external_source": rows[index].external_source,
                            "external_transaction_id": rows[index].external_transaction_id,
                        },
                        sort_keys=True,
                        ensure_ascii=False,
                    )
                    for index in indices
                    if rows[index].fields is not None
                }
                if len(semantics) > 1:
                    for index in indices:
                        rows[index] = replace(
                            rows[index], fields=None, issues=("IMPORT_SOURCE_ID_CONFLICT",)
                        )
            new_identifiers = set(new_ids)
            for index, row in enumerate(rows):
                if row.fields is None or row.fields.get("kind") != "expense_refund":
                    continue
                if row.fields.get("original_transaction_id") not in new_identifiers:
                    continue
                original_source_id = decoded[index].get("original_transaction_id", "")
                candidates = [rows[position] for position in groups.get(original_source_id, ())]
                original_row = next(
                    (
                        candidate
                        for candidate in candidates
                        if not candidate.issues
                        and candidate.fields is not None
                        and candidate.fields.get("kind") == "expense"
                    ),
                    None,
                )
                target = original_row.transaction_id if original_row else None
                if target is None:
                    for candidate in candidates:
                        existing = connection.execute(
                            "SELECT id FROM transactions WHERE kind='expense' "
                            "AND deleted_at_utc IS NULL AND external_source=? "
                            "AND external_transaction_id=?",
                            (candidate.external_source, candidate.external_transaction_id),
                        ).fetchone()
                        if existing:
                            target = str(existing[0])
                            break
                if target is None:
                    rows[index] = replace(
                        row, fields=None, issues=(*row.issues, "ORIGINAL_EXPENSE_REQUIRED")
                    )
                else:
                    rows[index] = replace(
                        row, fields={**row.fields, "original_transaction_id": target}
                    )
        preview = ImportPreview(
            table, mapping_json, mapping_hash, str(uuid4()), tuple(rows), revision
        )
        return replace(preview, plan_digest=_plan_digest(preview))

    def commit_payload(
        self, preview: ImportPreview, selected_rows: Iterable[int]
    ) -> dict[str, object]:
        """Recheck source bytes and selection; the ledger checks all live constraints."""
        if _plan_digest(preview) != preview.plan_digest:
            raise LedgerError("IMPORT_PREVIEW_CHANGED")
        if (
            hashlib.sha256(_source_bytes(preview.table.path)).hexdigest()
            != preview.table.file_digest
        ):
            raise LedgerError("IMPORT_FILE_CHANGED")
        selection = set(selected_rows)
        selected = [row for row in preview.rows if row.source_row_number in selection]
        if (
            not selected
            or len(selected) != len(selection)
            or any(row.issues or row.fields is None for row in selected)
        ):
            raise LedgerError("IMPORT_SELECTION_INVALID")
        available = {row.transaction_id for row in selected}
        all_ids = {row.transaction_id for row in preview.rows}
        for row in selected:
            original = cast(dict[str, object], row.fields).get("original_transaction_id")
            if original in all_ids and original not in available:
                raise LedgerError("IMPORT_DEPENDENCY_REQUIRED")
        return {
            "id": preview.batch_id,
            "source_format": preview.table.source_format,
            "source_file_name": preview.table.path.name,
            "file_digest": preview.table.file_digest,
            "mapping_hash": preview.mapping_hash,
            "mapping_json": preview.mapping_json,
            "format_version": 1,
            "source_row_count": len(preview.rows),
            "rows": tuple(
                {
                    "id": row.transaction_id,
                    "source_row_number": row.source_row_number,
                    "fields": row.fields,
                    "external_source": row.external_source,
                    "external_transaction_id": row.external_transaction_id,
                }
                for row in selected
            ),
        }

    def export_transactions(
        self,
        path: Path,
        format: str,
        filters: TransactionFilter = _DEFAULT_FILTER,
        *,
        cancel: Cancel = None,
    ) -> ExportResult:
        """Export exact decimal strings from one snapshot; default excludes deletions."""
        _validate(filters)
        where, parameters = _where(filters)
        with self.database.read() as connection:
            revision = _change_seq(connection)
            rows = tuple(
                dict(row)
                for row in connection.execute(
                    f"SELECT t.*,({_EFFECTIVE_BOOK}) AS effective_book_id,"
                    f"({_EFFECTIVE_CATEGORY}) AS effective_category_id,"
                    "book.name AS book_name,category.name AS category_name,"
                    "payment.name AS payment_method_name "
                    + _JOINS
                    + where
                    + "ORDER BY t.occurred_on,t.created_at_utc,t.id LIMIT ?",
                    [*parameters, MAX_ROWS + 1],
                )
            )
            if len(rows) > MAX_ROWS:
                raise LedgerError("EXPORT_TOO_LARGE")
            for offset in range(0, len(rows), 200):
                _cancelled(cancel)
                _decorate(connection, rows[offset : offset + 200])
            tags = {
                str(row[0]): str(row[1]) for row in connection.execute("SELECT id,name FROM tags")
            }
        values: list[tuple[str, ...]] = []
        for row in rows:
            _cancelled(cancel)
            row["transaction_id"] = row["id"]
            row["book_id"], row["category_id"] = (
                row["effective_book_id"],
                row["effective_category_id"],
            )
            row["format_version"] = "1"
            row["text_escape"] = "apostrophe-v1" if format == "csv" else "literal-v1"
            identifiers = cast(tuple[str, ...], row["tag_ids"])
            row["tag_names"] = json.dumps([tags[item] for item in identifiers], ensure_ascii=False)
            row["tag_ids"] = json.dumps(identifiers, ensure_ascii=False)
            values.append(
                tuple(str(row[key]) if row.get(key) is not None else "" for key in EXCHANGE_COLUMNS)
            )
        self._publish(path, format, EXCHANGE_COLUMNS, values, cancel=cancel, escaped_csv=True)
        return ExportResult(path.resolve(), len(rows), revision)

    def export_errors(self, preview: ImportPreview, path: Path, *, cancel: Cancel = None) -> Path:
        """Publish diagnostics without logging source contents."""
        rows = [
            (str(row.source_row_number), ";".join(row.issues), ";".join(row.possible_duplicates))
            for row in preview.rows
            if row.issues or row.possible_duplicates
        ]
        self._publish(
            path, "csv", ("source_row", "issues", "possible_duplicates"), rows, cancel=cancel
        )
        return path

    @staticmethod
    def _publish(
        path: Path,
        format: str,
        headers: tuple[str, ...],
        rows: Iterable[tuple[str, ...]],
        *,
        cancel: Cancel = None,
        escaped_csv: bool = False,
    ) -> None:
        if (
            format not in {"csv", "xlsx"}
            or path.suffix.lower() != "." + format
            or not path.parent.is_dir()
        ):
            raise LedgerError("EXPORT_PATH_INVALID")
        temporary: Path | None = None
        try:
            fd, name = tempfile.mkstemp(prefix=".openledger-", suffix="." + format, dir=path.parent)
            os.close(fd)
            temporary = Path(name)
            if format == "csv":
                with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(headers)
                    for row in rows:
                        _cancelled(cancel)
                        if escaped_csv:
                            writer.writerow(
                                tuple(
                                    value
                                    if key in {"format_version", "text_escape"}
                                    else "'" + value
                                    for key, value in zip(headers, row, strict=True)
                                )
                            )
                        else:
                            writer.writerow(
                                tuple(
                                    "'" + value
                                    if value.lstrip().startswith(("=", "+", "-", "@"))
                                    else value
                                    for value in row
                                )
                            )
                    handle.flush()
                    os.fsync(handle.fileno())
            else:
                workbook = Workbook()
                try:
                    sheet = workbook.active
                    if sheet is None:
                        raise LedgerError("EXPORT_IO_ERROR")
                    sheet.title = "OpenLedger"
                    sheet.append(headers)
                    for number, row in enumerate(rows, 2):
                        _cancelled(cancel)
                        for column, value in enumerate(row, 1):
                            cell = sheet.cell(number, column, value)
                            cell.data_type = "s"
                    sheet.freeze_panes = "A2"
                    sheet.auto_filter.ref = sheet.dimensions
                    workbook.save(temporary)
                finally:
                    workbook.close()
                with temporary.open("r+b") as handle:
                    os.fsync(handle.fileno())
            _cancelled(cancel)
            os.replace(temporary, path)
            temporary = None
        except OSError as error:
            raise LedgerError("EXPORT_IO_ERROR") from error
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
